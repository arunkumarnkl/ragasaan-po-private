"""Purchasing Word output and optional integrated alcohol inventory."""
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt
from po_logic import identifier, number
from po_v3_logic import pack_factor
from po_v34_features import purchase_category
from po_v35_features import inventory_items_for_source

def append_inventory_sheet(doc, alcohol_rows, order, ingredients):
    """Append all alcoholic PO items for this source, not just guide-selected rows.

    Includes beers, still/sparkling wines, whisky, vodka, gin, rum,
    tequila, cognac and liqueurs. Classification follows Ingredients.category.
    Every item is a separate, lined row; stock opens as ordered bottles.
    """
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    groups=inventory_items_for_source(order,ingredients)
    if not groups:
        return doc
    doc.add_page_break()
    doc.add_heading('Alcohol Inventory Sheet',0)
    doc.add_paragraph('Opening Stock = bottles supplied on this PO. Enter Unit Stock, Total, Consumption, Balance and Remarks after stock count.')
    headers=['Item','Opening Stock','Unit Stock','Total','Consumption','Balance','Remarks']
    widths=[2.18,.86,.72,.55,.80,.66,1.15]
    table=doc.add_table(rows=1,cols=len(headers))
    table.style='Table Grid'
    table.autofit=False
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for cell,heading,width in zip(table.rows[0].cells,headers,widths):
        cell.text=heading
        cell.width=Inches(width)
    for group,items in groups:
        row=table.add_row()
        row.cells[0].merge(row.cells[-1]).text=group.upper()
        for item in items:
            iid=identifier(item.get('ingredient_id'))
            master=ingredients.get(iid,{})
            supplied=number(item.get('Quantity'))*pack_factor(master,item.get('Source'))
            supplied_text=f'{supplied:g} bottles'
            cells=table.add_row().cells
            values=[item.get('Ingredient',''),supplied_text,'','','','','']
            for cell,value in zip(cells,values):cell.text=str(value)
            cells[1].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for row_index,row in enumerate(table.rows):
        row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height=Inches(.23)
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                for run in para.runs:
                    run.font.size=Pt(6.8)
                    if row_index==0:run.bold=True
                    if i==0 and row_index>0 and cell.text.isupper():run.bold=True
    return doc


def create_docx(order,metadata,logo_bytes=None,inventory_rows=None,all_orders=None,ingredients=None):
    """Compact source-specific PO. Guide formulas, pack sizes and prices remain sheet-driven."""
    from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import RGBColor
    doc=Document();section=doc.sections[0]
    section.top_margin=Inches(.55);section.bottom_margin=Inches(.55)
    section.left_margin=Inches(.62);section.right_margin=Inches(.62)
    # The logo and PURCHASE ORDER title share one compact, borderless header row.
    # Word's default paragraph layout puts consecutive paragraphs on separate lines.
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    header=doc.add_table(rows=1,cols=2)
    header.autofit=False
    header.columns[0].width=Inches(1.25)
    header.columns[1].width=Inches(5.85)
    logo_cell,title_cell=header.rows[0].cells
    logo_cell.width=Inches(1.25);title_cell.width=Inches(5.85)
    logo_cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    title_cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if logo_bytes:
        try:
            logo_p=logo_cell.paragraphs[0]
            logo_p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            logo_p.add_run().add_picture(BytesIO(logo_bytes),width=Inches(.88))
        except (OSError, ValueError, TypeError):pass
    title_p=title_cell.paragraphs[0]
    title_p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title_run=title_p.add_run('PURCHASE ORDER')
    title_run.bold=True;title_run.font.size=Pt(16)
    title_run.font.color.rgb=RGBColor(35,49,85)
    title_p.paragraph_format.space_after=Pt(2)
    is_unit=str(metadata.get('source','')).casefold()=='unit'
    details=[('Company',metadata.get('company','Ragasaan'),'Supplier',metadata.get('vendor','')),
             ('Party No', '__________' if is_unit else '', 'PO Date',metadata.get('po_date','')),
             ('Venue',metadata.get('venue',''),'Delivery',f"{metadata.get('delivery_date','')} {metadata.get('delivery_time','')}"),
             ('Address',metadata.get('address',''),'Contact',metadata.get('contact',''))]
    info=doc.add_table(rows=0,cols=4);info.autofit=False
    for desc in details:
        cells=info.add_row().cells
        for j,(cell,value,width) in enumerate(zip(cells,desc,[.78,2.4,.92,2.9])):
            cell.text=str(value or '');cell.width=Inches(width)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1)
                for run in para.runs:
                    run.font.size=Pt(7.3)
                    run.bold=j in (0,2)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    table=doc.add_table(rows=1,cols=8)
    # Unshaded, borderless PO body: no coloured background or grid lines.
    # The separately requested alcohol inventory retains its ruled grid.
    table.style=None;table.autofit=False
    headers=['ID','Item Name','Breakfast','Lunch','Dinner','Total','Unit','Remarks']
    # Compact numerical columns and minimal internal padding. More room for names.
    widths=[.46,2.90,.49,.43,.47,.50,.56,1.28]
    for i,width in enumerate(widths):
        table.columns[i].width=Inches(width)
    for i,(cell,name) in enumerate(zip(table.rows[0].cells,headers)):
        cell.text=name;cell.width=Inches(widths[i])
    for i in (2,3,4,5):
        header_para=table.rows[0].cells[i].paragraphs[0]
        header_para.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        for run in header_para.runs:run.font.size=Pt(6.5)
    for item in sorted(order,key=lambda x:(purchase_category(x.get('Category')),
                                              str(x['Ingredient']).casefold())):
        cells=table.add_row().cells
        cells[0].text=identifier(item.get('ingredient_id',''))
        cells[1].text=str(item['Ingredient'])
        breakdown=item.get('Meal display',{})
        for i,meal in enumerate(('Breakfast','Lunch','Dinner'),start=2):
            q=number(breakdown.get(meal,0))
            cells[i].text=f'{q:g}' if q else '—'
            para=cells[i].paragraphs[0]
            para.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            for run in para.runs:
                run.font.color.rgb=RGBColor(140,148,158)
                run.font.size=Pt(7)
        cells[5].text=f'{number(item["Quantity"]):g}'
        totalpara=cells[5].paragraphs[0]
        totalpara.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        for run in totalpara.runs:
            run.bold=True;run.font.size=Pt(8)
            run.font.color.rgb=RGBColor(22,32,50)
        cells[6].text=str(item['Unit'])
        cells[7].text=str(item['Remarks']).replace('Cocktail/Mocktail: ','')
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    def compact_cell_padding(cell, points=1.25):
        """Reduce the apparent gap between Breakfast and Lunch in Word."""
        tc_pr=cell._tc.get_or_add_tcPr()
        mar=tc_pr.first_child_found_in('w:tcMar')
        if mar is None:
            mar=OxmlElement('w:tcMar');tc_pr.append(mar)
        for edge in ('left','right'):
            node=mar.find(qn('w:'+edge))
            if node is None:
                node=OxmlElement('w:'+edge);mar.append(node)
            node.set(qn('w:w'),str(round(points*20)))
            node.set(qn('w:type'),'dxa')

    for row in table.rows:
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(.5)
                if i in (2,3,4,5):
                    compact_cell_padding(cell)
                for run in para.runs:
                    if run.font.size is None:run.font.size=Pt(7)
    if not is_unit:
        net=sum(number(x.get('Net')) for x in order)
        vat=number(metadata.get('vat',0))
        totals=doc.add_paragraph()
        totals.alignment=2
        text=totals.add_run(f'Net £{net:,.2f}    VAT ({vat:g}%) £{net*vat/100:,.2f}    Total £{net*(1+vat/100):,.2f}')
        text.font.size=Pt(7.5);text.font.color.rgb=RGBColor(156,163,175)
    if inventory_rows and all_orders is not None and ingredients is not None:
        append_inventory_sheet(doc,inventory_rows,all_orders,ingredients)
    buffer=BytesIO();doc.save(buffer);return buffer.getvalue()


def create_inventory_docx(alcohol_rows,order,ingredients):
    doc=Document()
    append_inventory_sheet(doc,alcohol_rows,order,ingredients)
    buffer=BytesIO();doc.save(buffer);return buffer.getvalue()


